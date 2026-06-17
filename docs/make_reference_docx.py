"""Build docs/reference.docx for pandoc's --reference-doc.

This reference document controls the look of the generated REPORT.docx:
  - sets the Normal and heading styles to Times New Roman, and
  - adds a centered PAGE-number field to the page footer.

Run from the repo root:  .venv/bin/python docs/make_reference_docx.py
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"


def set_style_font(document, style_name, size=None, bold=None):
    """Force a named style to use Times New Roman (and optional size/bold)."""
    try:
        style = document.styles[style_name]
    except KeyError:
        return
    font = style.font
    font.name = FONT
    # Make East-Asian / complex-script runs use the same face so the docx
    # does not silently fall back to Calibri.
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold


def add_page_number_field(paragraph):
    """Insert a Word PAGE field into the given paragraph."""
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    # A cached run value so the number shows before Word recalculates fields.
    cached = OxmlElement("w:r")
    cached_t = OxmlElement("w:t")
    cached_t.text = "1"
    cached.append(cached_t)
    fld.append(cached)
    run._r.addprevious(fld)


def main():
    doc = Document()

    # Base body font + common heading styles.
    set_style_font(doc, "Normal", size=12)
    for heading in (
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "TOC Heading",
    ):
        set_style_font(doc, heading)

    # Centered PAGE-number footer on the default section.
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Footer text should also be Times New Roman.
    para.style = doc.styles["Normal"]
    add_page_number_field(para)

    doc.save("docs/reference.docx")
    print("wrote docs/reference.docx")


if __name__ == "__main__":
    main()
